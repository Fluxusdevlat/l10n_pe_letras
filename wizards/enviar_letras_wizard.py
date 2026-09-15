from odoo import api, fields, models, _
from odoo.exceptions import UserError
import base64
from io import BytesIO
from datetime import date


class EnviarLetrasWizard(models.TransientModel):
    _name = 'l10n.pe.letras.email.wizard'
    _description = 'Enviar Letras por Email'

    date_from = fields.Date(string='Desde', required=True,
                            default=fields.Date.context_today)
    date_to = fields.Date(string='Hasta', required=True,
                          default=fields.Date.context_today)
    email_to = fields.Char(string='Enviar a', required=True)
    state = fields.Selection([
        ('draft', 'Borrador'),
        ('sent', 'Enviada'),
        ('signed', 'Firmada'),
        ('in_bank', 'En Banco'),
        ('protested', 'Protestada'),
        ('cancelled', 'Cancelada'),
        ('renewed', 'Renovada'),
        ('paid', 'Pagada'),
    ], string='Filtrar por Estado', default='in_bank')
    letra_count = fields.Integer(string='Letras a enviar',
                                 compute='_compute_letra_count')
    include_report = fields.Boolean(string='Incluir PDF por letra',
                                    default=True)
    include_excel = fields.Boolean(string='Incluir Excel resumen',
                                   default=True)

    @api.depends('date_from', 'date_to', 'state')
    def _compute_letra_count(self):
        for r in self:
            r.letra_count = self.env['l10n.pe.letra'].search_count([
                ('date_emission', '>=', r.date_from),
                ('date_emission', '<=', r.date_to),
                ('state', '=', r.state),
            ])

    @api.onchange('date_from', 'date_to')
    def _onchange_dates(self):
        if self.date_from and self.date_to and self.date_from > self.date_to:
            return {'warning': {
                'title': 'Fechas inválidas',
                'message': 'La fecha "Desde" no puede ser mayor a "Hasta".',
            }}

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        ICP = self.env['ir.config_parameter'].sudo()
        default_email = ICP.get_param('l10n_pe_letras.email_to', '')
        if default_email:
            res['email_to'] = default_email
        return res

    def action_preview(self):
        self.ensure_one()
        letras = self._get_letras()
        if not letras:
            raise UserError(_('No hay letras en el rango y estado seleccionados'))
        return {
            'type': 'ir.actions.act_window',
            'name': _('Letras a enviar (%s)') % len(letras),
            'res_model': 'l10n.pe.letra',
            'view_mode': 'list,form',
            'domain': [('id', 'in', letras.ids)],
        }

    def action_send(self):
        self.ensure_one()
        if not self.email_to:
            raise UserError(_('Debe especificar un correo destino'))

        letras = self._get_letras()
        if not letras:
            raise UserError(_('No hay letras en el rango y estado seleccionados'))

        today = date.today().strftime('%Y-%m-%d')
        attachment_ids = []

        # 1. Generar Excel resumen si aplica
        if self.include_excel:
            excel_data = self._generate_excel(letras)
            excel_attachment = self.env['ir.attachment'].sudo().create({
                'name': 'Resumen_Letras_%s.xlsx' % today,
                'type': 'binary',
                'datas': base64.b64encode(excel_data),
                'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            })
            attachment_ids.append(excel_attachment.id)

        # 2. Generar PDF por cada letra si aplica
        if self.include_report:
            report_action = self.env.ref('l10n_pe_letras.action_report_letra').sudo()
            for letra in letras:
                pdf, _ = report_action._render_qweb_pdf(letra.ids)
                pdf_attachment = self.env['ir.attachment'].sudo().create({
                    'name': 'Letra_%s.pdf' % letra.name.replace('/', '-'),
                    'type': 'binary',
                    'datas': base64.b64encode(pdf),
                    'mimetype': 'application/pdf',
                })
                attachment_ids.append(pdf_attachment.id)

        # 3. Cuerpo del correo
        lines = []
        for letra in letras:
            facturas = ', '.join(letra.line_ids.mapped('move_id.name'))
            lines.append(
                '• %s | %s | S/ %s | Vence: %s | Estado: %s | Facturas: %s' % (
                    letra.name, letra.partner_id.name,
                    '{:,.2f}'.format(letra.amount_total),
                    letra.date_due,
                    dict(letra._fields['state'].selection).get(letra.state),
                    facturas or 'Ninguna',
                )
            )

        body = _(
            '<h3>Reporte de Letras de Cambio</h3>'
            '<p>Período: %s al %s</p>'
            '<p>Estado: %s</p>'
            '<p>Total letras: <strong>%s</strong></p>'
            '<p>Monto total: <strong>S/ %s</strong></p>'
            '<hr/>'
            '<h4>Detalle:</h4>'
            '%s'
            '<hr/>'
            '<p style="color: #888;">Generado automáticamente por el sistema de Letras de Cambio</p>'
        ) % (
            self.date_from, self.date_to,
            dict(self._fields['state'].selection).get(self.state),
            len(letras),
            '{:,.2f}'.format(sum(letras.mapped('amount_total'))),
            '<br/>'.join(lines) if lines else '<em>Sin letras</em>',
        )

        # 4. Enviar correo real vía mail.mail
        subject = _('Letras de Cambio - %s al %s (%s letras)') % (
            self.date_from, self.date_to, len(letras))
        mail = self.env['mail.mail'].sudo().create({
            'subject': subject,
            'email_to': self.email_to,
            'body_html': body,
            'author_id': self.env.user.partner_id.id,
            'attachment_ids': [(6, 0, attachment_ids)],
        })
        mail.send()

        letras[0].message_post(
            body=body,
            subject=subject,
            attachment_ids=attachment_ids,
        )

        return {
            'type': 'ir.actions.act_window',
            'name': _('%s letras enviadas a %s') % (len(letras), self.email_to),
            'res_model': 'l10n.pe.letra',
            'view_mode': 'list,form',
            'domain': [('id', 'in', letras.ids)],
        }

    def _get_letras(self):
        self.ensure_one()
        return self.env['l10n.pe.letra'].search([
            ('date_emission', '>=', self.date_from),
            ('date_emission', '<=', self.date_to),
            ('state', '=', self.state),
        ])

    def _generate_excel(self, letras):
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = 'Letras de Cambio'

        # Encabezados
        headers = [
            'N° Letra', 'Cliente', 'RUC/DNI', 'Tipo', 'Fecha Emisión',
            'Fecha Vencimiento', 'Importe Total', 'Moneda', 'Estado',
            'Banco', 'Facturas Asociadas', 'Planilla', 'Observaciones'
        ]

        # Estilos
        header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        data_font = Font(name='Calibri', size=10)
        data_alignment = Alignment(vertical='center', wrap_text=False)
        money_format = '#,##0.00'

        # Escribir encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Escribir datos
        for row, letra in enumerate(letras, 2):
            facturas = ', '.join(letra.line_ids.mapped('move_id.name'))
            ruc = letra.partner_id.vat or ''
            estado = dict(letra._fields['state'].selection).get(letra.state, '')

            data = [
                letra.name,
                letra.partner_id.name,
                ruc,
                dict(letra._fields['tipo'].selection).get(letra.tipo, ''),
                letra.date_emission,
                letra.date_due,
                letra.amount_total,
                letra.currency_id.name,
                estado,
                letra.bank_id.name or '',
                facturas,
                letra.planilla_id.name or '',
                letra.notes or '',
            ]

            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                # Formato moneda en columna Importe Total (7)
                if col == 7:
                    cell.number_format = money_format

        # Anchos de columna
        col_widths = [14, 30, 15, 12, 14, 16, 14, 8, 14, 20, 30, 14, 30]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width

        # Fila de totales
        total_row = len(letras) + 2
        ws.cell(row=total_row, column=1, value='TOTALES').font = Font(bold=True, size=11)
        ws.cell(row=total_row, column=1).border = thin_border
        for col in range(2, 7):
            ws.cell(row=total_row, column=col).border = thin_border
        total_cell = ws.cell(row=total_row, column=7, value=sum(letras.mapped('amount_total')))
        total_cell.font = Font(bold=True, size=11)
        total_cell.number_format = money_format
        total_cell.border = thin_border
        for col in range(8, 14):
            ws.cell(row=total_row, column=col).border = thin_border

        # Información adicional al final
        info_row = total_row + 2
        ws.cell(row=info_row, column=1, value='Generado el:').font = Font(bold=True, size=10)
        ws.cell(row=info_row, column=2, value=fields.Date.today().strftime('%d/%m/%Y'))
        ws.cell(row=info_row + 1, column=1, value='Período:').font = Font(bold=True, size=10)
        ws.cell(row=info_row + 1, column=2, value='%s al %s' % (self.date_from, self.date_to))
        ws.cell(row=info_row + 2, column=1, value='Cantidad:').font = Font(bold=True, size=10)
        ws.cell(row=info_row + 2, column=2, value='%s letras' % len(letras))
        ws.cell(row=info_row + 3, column=1, value='Estado:').font = Font(bold=True, size=10)
        ws.cell(row=info_row + 3, column=2, value=dict(self._fields['state'].selection).get(self.state))

        # Configurar área de impresión
        ws.sheet_properties.pageSetUpPr = None
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1

        output = BytesIO()
        wb.save(output)
        return output.getvalue()
